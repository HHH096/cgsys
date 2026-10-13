# -*- coding: utf-8 -*-
"""生成「电气组本月行块」+ 对账数字，用于回填车队总表。

行块列序与车队总表 **七月下旬起的那套新表头**逐列一致（20 列）：

    物料名称/品牌 / 规格型号 / 申请数量 / 单位 / 预计单价 / 小计金额 / 运费 / 合计金额 /
    库存 / 用途 / 申购人 / 领用时间 / 申购人电话 / 样式/图片 / 网址链接 / 组别 / 备注 /
    实付金额 / 发票 / 是否到货

实测（`2026年赛车制造所需设备与耗材.xlsx`）：`七月下旬采购` / `八月采购` / `九月采购` /
`十月采购` 四个 sheet 都是这套 20 列；更早的月份是另外两套老表头（17 列、以及
`3月中旬采购` 那套带「请购数量/已购数量」的），**行块不能直接贴过去**。

**不含表头**：从第一行起就是数据，负责人在在线表格里「电气组那一块的末尾」直接粘贴即可。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app import config
from app.core import exporter, rules

BLOCK_COLUMNS = exporter.MAIN_COLUMNS
HEADERS = [c[1] for c in BLOCK_COLUMNS]
FIELD_ORDER = [c[0] for c in BLOCK_COLUMNS]

# 列顺序是照着总表定的，改这里等于改粘贴目标，务必和总表新表头保持一致
assert HEADERS == [
    '物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
    '合计金额', '库存', '用途', '申购人', '领用时间', '申购人电话', '样式/图片',
    '网址链接', '组别', '备注', '实付金额', '发票', '是否到货',
], '行块列序必须与车队总表的新表头一致'

HEADER_FILL = PatternFill('solid', fgColor='FF006600')
HEADER_FONT = Font(name='等线', size=11, bold=True, color='FFFFFFFF')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def fmt_num(value):
    """6200.0 → '6200'；5.10 → '5.1'；空值 → ''。"""
    if value is None or value == '':
        return ''
    if isinstance(value, bool):
        return str(value)
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value).strip()
    if abs(f - round(f)) < 1e-9:
        return str(int(round(f)))
    return ('%.2f' % f).rstrip('0').rstrip('.')


def cell_value(item, field):
    """取一条条目在某一列上的值。

    走 `rules.export_value`：退货那行的「实付金额」填净额、「备注」标「已退货」——
    与导出的三张 sheet 用同一套规则，免得两处不一致。
    """
    return rules.export_value(item, field)


def block_rows(items, group=None):
    """行块数据：list[list]，每行 20 列，顺序即 FIELD_ORDER。

    组别列在这里统一补上（条目自己填了就用它的，没填用传入的 group，再没有用默认组别），
    这样粘过去不用再手填。
    """
    group = (group or '').strip() or config.DEFAULT_GROUP
    rows = []
    for it in items:
        row = []
        for field in FIELD_ORDER:
            v = cell_value(it, field)
            if field == 'group_name':
                v = (v or '').strip() or group
            row.append(v)
        rows.append(row)
    return rows


def block_tsv(items, group=None):
    """行块 TSV 文本（制表符分隔），贴到在线表格里会自己按列展开。"""
    lines = []
    for row in block_rows(items, group):
        lines.append('\t'.join(fmt_num(v) for v in row))
    return '\n'.join(lines)


def reconcile(items, round_=None):
    """对账数字：预计 / 实付 / 差额 / 运费 / 没填实付的条数 / 各商家小计。"""
    items = list(items or [])
    estimate = rules.r2(sum(rules.to_num(it.get('total')) or 0.0 for it in items)) or 0.0
    freight = rules.r2(sum(rules.to_num(it.get('freight')) or 0.0 for it in items)) or 0.0

    paid_vals = [rules.to_num(it.get('paid')) for it in items]
    paid_vals = [v for v in paid_vals if v is not None]
    paid = rules.r2(sum(paid_vals)) or 0.0

    # 退货：单独看退回来多少、真正净花了多少
    refunded = [it for it in items if it.get('refunded')]
    refund = rules.r2(sum(rules.to_num(it.get('refund_amount')) or 0.0
                          for it in refunded)) or 0.0
    net = rules.r2(paid - refund) or 0.0

    _, _, group_totals = rules.split_items(
        items, (round_ or {}).get('split_threshold'))
    suppliers = [{'supplier': k.split(':', 1)[1] or '(未填商家)', 'sum': v}
                 for k, v in group_totals.items()]

    return {
        'count': len(items),
        'estimate': estimate,
        'paid': paid,
        'refund': refund,
        'net': net,
        'refunded_count': len(refunded),
        # 差额按**净支出**算 —— 退回来的钱本来就不算花掉的
        'diff': rules.r2(net - estimate) or 0.0,
        'freight': freight,
        'paid_count': len(paid_vals),
        'unpriced': len(items) - len(paid_vals),
        'suppliers': suppliers,
    }


def build_block_workbook(items, group=None, recon=None):
    """返回 Workbook：`行块`（纯数据，从 A1 起）+ `对账`。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '行块'
    for row in block_rows(items, group):
        ws.append([v if v != '' else None for v in row])
    for letter in ('ABCDEFGHIJKLMNOPQRST')[:len(FIELD_ORDER)]:
        ws.column_dimensions[letter].width = 14

    rs = wb.create_sheet('对账')
    rs.append(['项目', '数值'])
    rs['A1'].fill = HEADER_FILL
    rs['B1'].fill = HEADER_FILL
    rs['A1'].font = HEADER_FONT
    rs['B1'].font = HEADER_FONT
    recon = recon or reconcile(items)
    for label, value in (
            ('条目数', recon['count']),
            ('预计合计（合计金额之和）', recon['estimate']),
            ('实付合计', recon['paid']),
            ('其中退货退款合计', recon['refund']),
            ('净支出（实付 − 退款）', recon['net']),
            ('差额（净支出 − 预计）', recon['diff']),
            ('其中运费合计', recon['freight']),
            ('已填实付的条数', recon['paid_count']),
            ('没填实付的条数', recon['unpriced']),
            ('已退货的条数', recon['refunded_count']),
    ):
        rs.append([label, value])
    rs.append([])
    rs.append(['按商家小计', ''])
    for s in recon['suppliers']:
        rs.append([s['supplier'], s['sum']])
    rs.column_dimensions['A'].width = 26
    rs.column_dimensions['B'].width = 14
    return wb


def save_block_xlsx(items, path, group=None, recon=None):
    wb = build_block_workbook(items, group, recon)
    wb.save(str(path))
    return path
