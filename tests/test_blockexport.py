# -*- coding: utf-8 -*-
"""P4 行块生成与对账数字。"""
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app import config
from app.core import blockexport, repo
from tests.support import TempDataTest

MASTER = Path(r'C:\Users\<你的用户名>\Desktop\2026年赛车制造所需设备与耗材.xlsx')
NEW_HEADER_SHEETS = ['7月下旬采购', '八月采购', '九月采购', '十月采购']

EXPECTED_HEADERS = [
    '物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
    '合计金额', '库存', '用途', '申购人', '领用时间', '申购人电话', '样式/图片',
    '网址链接', '组别', '备注', '实付金额', '发票', '是否到货',
]


def load_master():
    """总表里有个 openpyxl 读不了的空 <fill/>，打个补丁再读。"""
    try:
        zin = zipfile.ZipFile(str(MASTER))
    except (OSError, zipfile.BadZipFile) as exc:
        # 桌面那份表可能被挪走、改名，或者是 OneDrive 的"仅在线"占位符 ——
        # **存在但打不开**。这种机器上核对不了列序，跳过就好；
        # 报 error 会让人以为是自己刚改坏了代码（真发生过）。
        raise unittest.SkipTest('车队总表打不开：%s（%s）' % (MASTER, exc))
    tmp = os.path.join(tempfile.gettempdir(), 'master_for_block_test.xlsx')
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'xl/styles.xml':
                data = data.decode('utf-8').replace(
                    '<fill/>', '<fill><patternFill patternType="none"/></fill>').encode('utf-8')
            zout.writestr(item, data)
    return load_workbook(tmp, data_only=True)


@unittest.skipUnless(MASTER.exists(), '需要车队总表文件才能核对列序')
class TestMatchesRealMaster(unittest.TestCase):
    """列序必须和车队总表的新表头逐列一致，否则粘过去整行错位。"""

    @classmethod
    def setUpClass(cls):
        cls.wb = load_master()

    def test_every_new_header_sheet_matches(self):
        for sheet in NEW_HEADER_SHEETS:
            ws = self.wb[sheet]
            got = [str(ws.cell(row=2, column=c).value).strip()
                   for c in range(1, len(EXPECTED_HEADERS) + 1)]
            self.assertEqual(got, EXPECTED_HEADERS, '%s 的表头与行块不一致' % sheet)

    def test_block_headers_equal_that_header(self):
        self.assertEqual(blockexport.HEADERS, EXPECTED_HEADERS)

    def test_old_sheets_really_use_a_different_header(self):
        """老月份是另一套表头 —— 行块贴过去会错位，界面上要提醒。"""
        ws = self.wb['三月下旬采购'] if '三月下旬采购' in self.wb.sheetnames else self.wb['3月下旬采购']
        got = [str(ws.cell(row=2, column=c).value).strip()
               for c in range(1, len(EXPECTED_HEADERS) + 1)]
        self.assertNotEqual(got, EXPECTED_HEADERS)
        self.assertIn('预计金额', got, '老表头里应当有「预计金额」这一列')


class TestFmt(unittest.TestCase):

    def test_integers_lose_the_decimal(self):
        self.assertEqual(blockexport.fmt_num(6200.0), '6200')
        self.assertEqual(blockexport.fmt_num(2), '2')

    def test_decimals_keep_significant_digits(self):
        self.assertEqual(blockexport.fmt_num(5.10), '5.1')
        self.assertEqual(blockexport.fmt_num(8.67), '8.67')
        self.assertEqual(blockexport.fmt_num(157.92), '157.92')

    def test_empty_values(self):
        self.assertEqual(blockexport.fmt_num(None), '')
        self.assertEqual(blockexport.fmt_num(''), '')

    def test_text_passes_through(self):
        self.assertEqual(blockexport.fmt_num('不锈钢螺丝'), '不锈钢螺丝')


class TestBlockRows(unittest.TestCase):

    def items(self):
        return [
            {'name': '充电机', 'spec': '3.3kwh', 'qty': 1, 'unit': '个', 'price': 6200,
             'subtotal': 6200, 'freight': 0, 'total': 6200, 'stock': '', 'purpose': '高压电池充电',
             'applicant': '张三', 'used_at': '', 'phone': '13800000000', 'image_path': '',
             'url': '浩泰科技', 'group_name': '电气组', 'note': '', 'paid': 6200,
             'invoice': '有', 'arrived': '是'},
            {'name': '绝缘手套', 'spec': '/', 'qty': 2, 'unit': '个', 'price': 250,
             'subtotal': 500, 'freight': 0, 'total': 500, 'purpose': '箱检',
             'applicant': '李四', 'phone': '15100000000', 'url': '速迫',
             'group_name': '电气组', 'paid': None},
        ]

    def test_shape(self):
        rows = blockexport.block_rows(self.items())
        self.assertEqual(len(rows), 2)
        self.assertEqual([len(r) for r in rows], [20, 20])

    def test_values_line_up_with_columns(self):
        row = blockexport.block_rows(self.items())[0]
        self.assertEqual(row[0], '充电机')                      # 物料名称/品牌
        self.assertEqual(row[1], '3.3kwh')                     # 规格型号
        self.assertEqual(row[2], 1)                            # 申请数量
        self.assertEqual(row[3], '个')                          # 单位
        self.assertEqual(row[4], 6200)                         # 预计单价
        self.assertEqual(row[5], 6200)                         # 小计金额
        self.assertEqual(row[6], 0)                            # 运费
        self.assertEqual(row[7], 6200)                         # 合计金额
        self.assertEqual(row[9], '高压电池充电')                 # 用途
        self.assertEqual(row[10], '张三')                     # 申购人
        self.assertEqual(row[12], '13800000000')               # 申购人电话
        self.assertEqual(row[14], '浩泰科技')                    # 网址链接
        self.assertEqual(row[15], '电气组')                      # 组别
        self.assertEqual(row[17], 6200)                        # 实付金额
        self.assertEqual(row[18], '有')                         # 发票
        self.assertEqual(row[19], '是')                         # 是否到货

    def test_group_is_filled_even_when_missing(self):
        rows = blockexport.block_rows([{'name': 'A', 'group_name': ''}])
        self.assertEqual(rows[0][15], '电气组')

    def test_group_can_be_overridden(self):
        rows = blockexport.block_rows([{'name': 'A', 'group_name': ''}], group='动力组')
        self.assertEqual(rows[0][15], '动力组')

    def test_item_own_group_wins_over_the_override(self):
        rows = blockexport.block_rows([{'name': 'A', 'group_name': '动力组'}], group='电气组')
        self.assertEqual(rows[0][15], '动力组')

    def test_blank_item_does_not_crash(self):
        rows = blockexport.block_rows([{}])
        self.assertEqual(len(rows[0]), 20)
        self.assertIsNone(rows[0][0], '空条目的名称列应当是 None（对应 Excel 空单元格）')
        self.assertEqual(rows[0][15], '电气组')


class TestTsv(unittest.TestCase):

    def test_one_line_per_item_and_no_header(self):
        items = [{'name': 'A'}, {'name': 'B'}]
        text = blockexport.block_tsv(items)
        lines = text.split('\n')
        self.assertEqual(len(lines), 2, '不得包含表头行')
        self.assertTrue(lines[0].startswith('A\t'))
        self.assertNotIn('物料名称/品牌', text)

    def test_columns_are_tab_separated(self):
        text = blockexport.block_tsv([{'name': '充电机', 'unit': '个', 'qty': 1}])
        self.assertEqual(len(text.split('\t')), 20)

    def test_numbers_are_written_cleanly(self):
        text = blockexport.block_tsv([{'name': 'A', 'total': 6200.0, 'qty': 2.0}])
        cells = text.split('\t')
        self.assertEqual(cells[7], '6200')      # 合计金额
        self.assertEqual(cells[2], '2')         # 申请数量

    def test_empty_items_gives_empty_text(self):
        self.assertEqual(blockexport.block_tsv([]), '')


class TestReconcile(unittest.TestCase):

    def items(self):
        return [
            {'name': 'A', 'total': 6200, 'paid': 6200, 'freight': 0, 'url': '浩泰科技'},
            {'name': 'B', 'total': 1050, 'paid': 1030, 'freight': 0, 'url': '诺可得'},
            {'name': 'C', 'total': 500, 'paid': None, 'freight': 30, 'url': '速迫'},
        ]

    def test_totals(self):
        r = blockexport.reconcile(self.items())
        self.assertEqual(r['count'], 3)
        self.assertEqual(r['estimate'], 7750.0)
        self.assertEqual(r['paid'], 7230.0)
        self.assertEqual(r['diff'], -520.0)
        self.assertEqual(r['freight'], 30.0)

    def test_counts_unpriced_rows(self):
        r = blockexport.reconcile(self.items())
        self.assertEqual(r['paid_count'], 2)
        self.assertEqual(r['unpriced'], 1, '没填实付的条数要报出来，否则差额是虚的')

    def test_suppliers_are_listed(self):
        r = blockexport.reconcile(self.items())
        names = [s['supplier'] for s in r['suppliers']]
        self.assertEqual(names, ['浩泰科技', '诺可得', '速迫'])

    def test_no_money_float_noise(self):
        r = blockexport.reconcile([{'total': 0.1, 'paid': 0.2}, {'total': 0.2, 'paid': 0.4}])
        self.assertEqual(r['estimate'], 0.3)
        self.assertEqual(r['paid'], 0.6)
        self.assertEqual(r['diff'], 0.3)

    def test_empty(self):
        r = blockexport.reconcile([])
        self.assertEqual((r['count'], r['estimate'], r['paid'], r['diff']), (0, 0.0, 0.0, 0.0))


class TestBlockWorkbook(TempDataTest):

    def setUp(self):
        super().setUp()
        self.mid = repo.add_member('张三', '13800000000')
        self.rid = repo.create_round('八月采购')
        repo.add_item(self.rid, self.mid, {
            'name': '充电机', 'spec': '3.3kwh', 'qty': 1, 'unit': '个', 'price': 6200,
            'freight': 0, 'purpose': '高压电池充电', 'applicant': '张三',
            'url': '浩泰科技', 'paid': 6200})
        repo.add_item(self.rid, self.mid, {
            'name': '绝缘手套', 'qty': 2, 'unit': '个', 'price': 250,
            'purpose': '箱检', 'applicant': '李四', 'url': '速迫'})
        self.items = repo.list_items(self.rid)

    def save(self):
        path = config.EXPORT_DIR / 'block.xlsx'
        config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        blockexport.save_block_xlsx(self.items, path)
        return load_workbook(str(path))

    def test_sheet_names(self):
        self.assertEqual(self.save().sheetnames, ['行块', '对账'])

    def test_block_sheet_starts_with_data_immediately(self):
        ws = self.save()['行块']
        self.assertEqual(ws.cell(row=1, column=1).value, '充电机',
                         '第一行就该是数据，不能是表头 —— 否则粘贴时会多一行')
        self.assertEqual(ws.cell(row=2, column=1).value, '绝缘手套')
        self.assertEqual(ws.cell(row=1, column=16).value, '电气组')

    def test_reconcile_sheet_has_the_numbers(self):
        ws = self.save()['对账']
        rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                for r in range(1, ws.max_row + 1)}
        self.assertEqual(rows['条目数'], 2)
        self.assertEqual(rows['预计合计（合计金额之和）'], 6700)
        self.assertEqual(rows['实付合计'], 6200)
        self.assertEqual(rows['没填实付的条数'], 1)

    def test_item_values_are_written(self):
        ws = self.save()['行块']
        self.assertEqual(ws.cell(row=1, column=8).value, 6200)    # 合计金额
        self.assertEqual(ws.cell(row=1, column=15).value, '浩泰科技')
        self.assertEqual(ws.cell(row=2, column=8).value, 500)


if __name__ == '__main__':
    unittest.main(verbosity=2)
