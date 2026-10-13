# -*- coding: utf-8 -*-
"""端到端验收：起真实 HTTP 服务，走完整流程。

不是直接调函数，而是真的起一个 uvicorn 服务、用 httpx 跑完整流程，
连导出的 xlsx 与上传的发票文件都按真实落盘情况核对。
"""
import io
import socket
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

import httpx
import uvicorn
from openpyxl import load_workbook

from app import config
from app.server.api import app
from tests.support import temp_config_patches

MAIN_HEADERS = ['物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
                '合计金额', '库存', '用途', '申购人', '领用时间', '申购人电话', '样式/图片',
                '网址链接', '组别', '备注', '实付金额', '发票', '是否到货']

SPLIT_HEADERS = ['物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
                 '合计金额', '库存', '用途', '申购人', '领用时间', '实际金额', '申购人电话',
                 '样式/图片', '网址链接', '组别', '备注', '实际金额', '发票']

PDF = b'%PDF-1.4 fake invoice body'


def free_port():
    s = socket.socket()
    s.bind(('127.0.0.1', 0))
    port = s.getsockname()[1]
    s.close()
    return port


class LiveServerCase(unittest.TestCase):
    """起一个真实的 uvicorn 服务，数据目录指到临时目录。

    注意：配置必须在起服务**之前**就打补丁 —— 服务的 lifespan 会建库，
    否则会在真实的 data 目录里建库。
    """

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls._patches = []
        for name, val in temp_config_patches(cls._tmp.name):
            p = mock.patch.object(config, name, val)
            p.start()
            cls._patches.append(p)

        cls.port = free_port()
        cls.config = uvicorn.Config(app, host='127.0.0.1', port=cls.port,
                                    log_level='error', lifespan='on')
        cls.server = uvicorn.Server(cls.config)
        cls.thread = threading.Thread(target=cls.server.run, daemon=True)
        cls.thread.start()
        for _ in range(200):
            if cls.server.started:
                break
            threading.Event().wait(0.05)
        if not cls.server.started:
            raise RuntimeError('服务没能在 10 秒内起来')

    @classmethod
    def tearDownClass(cls):
        cls.server.should_exit = True
        cls.thread.join(timeout=10)
        for p in cls._patches:
            p.stop()
        cls._tmp.cleanup()

    def setUp(self):
        self.client = httpx.Client(base_url='http://127.0.0.1:%d' % self.port, timeout=20)
        self.login_purchase()

    def tearDown(self):
        self.client.close()

    def login_purchase(self, password='test-pw'):
        """接口现在分「采购人员」和「组员」两种身份，采购专用接口必须登录。

        第一次是「还没设过密码」，可以直接设；之后就要走登录接口了。
        """
        st = self.client.get('/api/purchase/status').json()
        if not st['has_password']:
            r = self.client.post('/api/purchase/password', json={'new': password})
        else:
            r = self.client.post('/api/purchase/login', json={'password': password})
        self.assertEqual(r.status_code, 200, r.text)
        return password

    # ------------------------------------------------------------ 小工具

    def new_round(self, name='八月采购'):
        r = self.client.post('/api/rounds', json={'name': name})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['id']

    def new_member(self, name, phone='', is_owner=False):
        r = self.client.post('/api/members', json={'name': name, 'phone': phone,
                                                   'is_owner': is_owner})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['id']

    def add_item(self, rid, mid, **kw):
        body = {'member_id': mid, 'applicant': '张三', 'group_name': '电气组',
                'unit': '个', 'qty': 1, 'purpose': '用', 'price': 1, 'freight': 0}
        body.update(kw)
        r = self.client.post('/api/rounds/%d/items' % rid, json=body)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['item']

    def set_status(self, iid, **kw):
        r = self.client.post('/api/items/%d/status' % iid, json=kw)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def upload(self, iid, name='invoice.pdf', data=PDF):
        return self.client.post('/api/items/%d/invoices' % iid,
                                files={'file': (name, io.BytesIO(data), 'application/pdf')})

    def detail(self, rid):
        r = self.client.get('/api/rounds/%d' % rid)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def export_wb(self, rid):
        r = self.client.get('/api/rounds/%d/export' % rid)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn('spreadsheetml', r.headers['content-type'])
        self.assertEqual(r.content[:2], b'PK')
        return load_workbook(io.BytesIO(r.content)), r


# ====================================================================== P1

# 商家 -> 期望落在哪个分表（验证「按同一商家本轮合计」这条规则）
ROWS = [
    dict(name='充电机', spec='3.3kwh', qty=1, unit='个', price=6200, freight=0,
         purpose='高压电池充电', url='浩泰科技', expect='high'),
    dict(name='绝缘手套', spec='/', qty=1, unit='个', price=300, freight=0,
         purpose='箱检', url='速迫', expect='low'),
    dict(name='绝缘鞋', spec='/', qty=1, unit='个', price=300, freight=0,
         purpose='箱检', url='速迫', expect='low'),
    dict(name='强芯毡', spec='2mm*1m', qty=1, unit='张', price=600, freight=0,
         purpose='制作空套', url='淘宝商家', expect='high'),
    dict(name='蜂窝铝', spec='5*1000*2000', qty=1, unit='块', price=600, freight=30,
         purpose='制作空套', url='淘宝商家', expect='high'),
    dict(name='螺丝收纳盒', spec='14内盒', qty=2, unit='个', price=400, freight=0,
         purpose='收纳', url='', expect='low'),
]


class TestP1RoundTrip(LiveServerCase):

    def test_full_round_trip(self):
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('采购系统', r.text)
        self.assertEqual(self.client.get('/app.js').status_code, 200)

        self.assertEqual(self.client.get('/api/members').json(), [])
        mid = self.new_member('张三', '13800000000', is_owner=True)
        self.new_member('李四', '15100000000')

        rid = self.new_round()
        self.assertEqual(self.client.get('/api/rounds/%d' % rid).json()['round']['status'], 'open')

        for i, row in enumerate(ROWS):
            body = dict(row)
            body.pop('expect')
            it = self.add_item(rid, mid, **body)
            self.assertEqual(it['subtotal'], round(body['qty'] * body['price'], 2))
            self.assertEqual(it['total'], round(body['qty'] * body['price'] + body['freight'], 2))

        split = self.client.get('/api/rounds/%d/preview' % rid).json()
        self.assertEqual(split['high_count'], sum(1 for r in ROWS if r['expect'] == 'high'))
        self.assertEqual(split['low_count'], sum(1 for r in ROWS if r['expect'] == 'low'))
        self.assertEqual(split['all_sum'],
                         round(sum(r['qty'] * r['price'] + r['freight'] for r in ROWS), 2))
        self.assertEqual(split['threshold'], 1000.0)

        wb, r = self.export_wb(rid)
        self.assertTrue(r.headers.get('content-disposition'))
        self.assertEqual(wb.sheetnames, ['总表', '低于一千', '高于一千'])

        main = wb['总表']
        self.assertEqual([main.cell(row=2, column=c).value for c in range(1, 21)], MAIN_HEADERS)
        self.assertIsNone(main.cell(row=2, column=21).value)
        for sheet, last in (('低于一千', '是否到货'), ('高于一千', '是否有合同')):
            ws = wb[sheet]
            self.assertEqual([ws.cell(row=2, column=c).value for c in range(1, 22)],
                             SPLIT_HEADERS + [last], sheet)

        self.assertEqual([main.cell(row=r, column=1).value for r in range(3, 3 + len(ROWS))],
                         [r['name'] for r in ROWS])
        self.assertEqual([wb['低于一千'].cell(row=r, column=1).value
                          for r in range(3, 3 + split['low_count'])],
                         [r['name'] for r in ROWS if r['expect'] == 'low'])
        self.assertEqual([wb['高于一千'].cell(row=r, column=1).value
                          for r in range(3, 3 + split['high_count'])],
                         [r['name'] for r in ROWS if r['expect'] == 'high'])

        r_total = 3 + len(ROWS)
        self.assertEqual(main.cell(row=r_total, column=1).value, '合计')
        self.assertEqual(main.cell(row=r_total, column=8).value, '=SUM(H3:H%d)' % (r_total - 1))
        low, r_low = wb['低于一千'], 3 + split['low_count']
        self.assertEqual(low.cell(row=r_low, column=7).value, '合计')
        self.assertIsNone(low.cell(row=r_low, column=1).value)
        high, r_high = wb['高于一千'], 3 + split['high_count']
        self.assertEqual(high.cell(row=r_high, column=1).value, '合计')
        self.assertIn('A%d:G%d' % (r_high, r_high), [str(x) for x in high.merged_cells.ranges])

        def col_sum(ws, n):
            return round(sum(ws.cell(row=r, column=8).value or 0 for r in range(3, 3 + n)), 2)

        self.assertEqual(col_sum(low, split['low_count']) + col_sum(high, split['high_count']),
                         col_sum(main, len(ROWS)))

        r = self.client.post('/api/rounds/%d/status' % rid, json={'status': 'locked'})
        self.assertEqual(r.json()['status'], 'locked')
        body = dict(ROWS[0])
        body.pop('expect')
        body.update(member_id=mid, applicant='张三', unit='个', phone='')
        self.assertEqual(self.client.post('/api/rounds/%d/items' % rid, json=body).status_code, 400)
        self.assertEqual(self.client.get('/api/rounds/%d/export' % rid).status_code, 200)


# ====================================================================== P2

class TestP2ProcurementFlow(LiveServerCase):
    """下单 -> 传发票 -> 重命名归档 -> 缺发票清单 -> 导出。"""

    def make_round(self):
        mid = self.new_member('张三', '13800000000', is_owner=True)
        rid = self.new_round()
        items = [
            self.add_item(rid, mid, name='充电机', spec='3.3kwh', unit='个', price=6200,
                          purpose='高压电池充电', url='浩泰科技', paid=6200),
            self.add_item(rid, mid, name='绝缘毯', spec='1000*1000', unit='张', qty=2,
                          price=1050, purpose='高压保护', url='商家诺可得商家', paid=2060),
            self.add_item(rid, mid, name='绝缘手套', spec='/', unit='个', qty=2,
                          price=250, purpose='箱检', url='速迫', paid=500),
            self.add_item(rid, mid, name='钮子开关', spec='9011两脚两挡', unit='个', qty=2,
                          price=23, purpose='挡位开关', url='https://e.tb.cn/h.8VOOpKHaRc',
                          note='张勇已付款'),
        ]
        return rid, mid, items

    def test_ordering_then_invoicing_then_export(self):
        rid, mid, items = self.make_round()
        ids = [i['id'] for i in items]

        # 1) 一上来四条都待下单、都缺发票
        d = self.detail(rid)
        self.assertEqual(d['pending_order'], ids)
        self.assertEqual(d['pending_invoice'], ids)
        self.assertEqual(d['split_ids']['high'], [ids[0], ids[1]])
        self.assertEqual(d['split_ids']['low'], [ids[2], ids[3]])

        # 2) 逐条下单，待下单清单随之编短
        for i, iid in enumerate(ids):
            self.set_status(iid, order_status='已下单')
            self.assertEqual(self.detail(rid)['pending_order'], ids[i + 1:])

        # 3) 到货一条
        self.set_status(ids[0], order_status='已到货')
        d = self.detail(rid)
        self.assertEqual([i['id'] for i in d['items'] if i['order_status'] == '已到货'], [ids[0]])

        # 4) 传真发票：重命名 + 落盘
        r = self.upload(ids[0], name='任意名字.pdf')
        self.assertEqual(r.status_code, 200, r.text)
        up = r.json()
        self.assertEqual(up['state'], 'new')
        self.assertEqual(up['invoice']['stored_name'], 'DQ-充电机-6200.pdf')
        self.assertEqual(up['invoice']['rel_path'], 'invoices/八月采购/DQ-充电机-6200.pdf')
        archived = config.DATA_DIR / up['invoice']['rel_path']
        self.assertTrue(archived.exists(), '发票没有真的落到磁盘')
        self.assertEqual(archived.read_bytes(), PDF)

        # 5) 收了发票就自动离开缺发票清单
        self.assertEqual(self.detail(rid)['pending_invoice'], ids[1:])
        self.assertEqual(up['item']['invoice_status'], '已收到')
        self.assertEqual(up['item']['invoice'], '有')

        # 6) 同一份文件再传一次：不产生副本
        again = self.upload(ids[0], name='又传一次.pdf').json()
        self.assertEqual(again['state'], 'duplicate')
        self.assertEqual(len(list((config.INVOICE_DIR / '八月采购').iterdir())), 1)

        # 7) 再给一条传发票，然后导出
        self.upload(ids[1], name='bill.pdf')
        wb, _ = self.export_wb(rid)

        # 8) 导出里「发票」「是否到货」两列跟着状态走
        main = wb['总表']
        self.assertEqual(main.cell(row=3, column=1).value, '充电机')
        self.assertEqual(main.cell(row=3, column=19).value, '有')     # 发票
        self.assertEqual(main.cell(row=3, column=20).value, '是')     # 是否到货
        self.assertEqual(main.cell(row=4, column=1).value, '绝缘毯')
        self.assertEqual(main.cell(row=4, column=19).value, '有')
        self.assertIsNone(main.cell(row=4, column=20).value)          # 还没到货
        self.assertIsNone(main.cell(row=5, column=19).value)          # 绝缘手套还没发票
        self.assertIsNone(main.cell(row=6, column=19).value)          # 钮子开关还没发票

        # 9) 分表也带着这两列
        high = wb['高于一千']
        self.assertEqual([high.cell(row=r, column=1).value for r in (3, 4)], ['充电机', '绝缘毯'])
        self.assertEqual(high.cell(row=3, column=20).value, '有')
        low = wb['低于一千']
        self.assertEqual([low.cell(row=r, column=1).value for r in (3, 4)], ['绝缘手套', '钮子开关'])

        # 10) 删掉一张发票：记录、文件、清单三处都要跟上
        inv_id = up['invoice']['id']
        self.assertEqual(self.client.delete('/api/invoices/%d' % inv_id).status_code, 200)
        self.assertFalse(archived.exists(), '删记录时文件应一起删')
        d = self.detail(rid)
        self.assertEqual(d['pending_invoice'], [ids[0], ids[2], ids[3]])
        self.assertEqual(d['invoice_counts'], {str(ids[1]): 1})

        # 11) 锁定之后：改条目被拒，但下单与收发票照常
        self.client.post('/api/rounds/%d/status' % rid, json={'status': 'locked'})
        self.assertEqual(self.client.post('/api/rounds/%d/items' % rid, json={
            'member_id': mid, 'name': 'x', 'unit': '个', 'qty': 1, 'price': 1,
            'purpose': 'y', 'applicant': '张三'}).status_code, 400)
        self.assertEqual(self.set_status(ids[2], order_status='已到货')['order_status'], '已到货')
        self.assertEqual(self.upload(ids[2], name='later.pdf').status_code, 200)
        self.assertEqual(self.detail(rid)['pending_invoice'], [ids[0], ids[3]])

        # 12) 锁定后导出照样能拿到，且反映最新状态
        wb2, _ = self.export_wb(rid)
        self.assertEqual(wb2['总表'].cell(row=5, column=19).value, '有')
        self.assertEqual(wb2['总表'].cell(row=5, column=20).value, '是')


if __name__ == '__main__':
    unittest.main(verbosity=2)
