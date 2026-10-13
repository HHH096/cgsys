# -*- coding: utf-8 -*-
"""P4 端到端验收：把行块真的「贴」进车队总表副本，再按总表表头逐列读回来核对。

这是「贴过去不会错位」的证明 —— 不是比字符串，而是：
  起真实服务 → 建轮次加条目 → 取行块 TSV → 写进总表副本 →
  用**总表自己的表头**建立「列名 → 列号」映射 → 逐列读回验证语义正确。

绝不修改你的总表原文件：全程操作的是临时副本。
"""
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from tests.test_e2e import LiveServerCase

MASTER = Path(r'C:\Users\<你的用户名>\Desktop\2026年赛车制造所需设备与耗材.xlsx')
TARGET_SHEET = '十月采购'          # 一个用新表头的月份


def patched_master_copy(dest):
    """总表里有 openpyxl 读不了的 styles.xml（空 <fill/>），打补丁拷一份到 dest。"""
    try:
        zin = zipfile.ZipFile(str(MASTER))
    except (OSError, zipfile.BadZipFile) as exc:
        # 同上：桌面那份表可能被挪走/改名，或者是 OneDrive 的"仅在线"占位符
        # （存在但打不开）。核对不了就跳过，别让整个套件报红。
        raise unittest.SkipTest('车队总表打不开：%s（%s）' % (MASTER, exc))
    with zipfile.ZipFile(str(dest), 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'xl/styles.xml':
                data = data.decode('utf-8').replace(
                    '<fill/>', '<fill><patternFill patternType="none"/></fill>').encode('utf-8')
            zout.writestr(item, data)


def header_map(ws, row=2):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=row, column=c).value
        if v is not None:
            out[str(v).strip()] = c
    return out


def coerce(val):
    """像电子表格粘贴那样：数字变数字，其余保持文本。"""
    try:
        f = float(val)
    except (TypeError, ValueError):
        return val
    return int(f) if abs(f - round(f)) < 1e-9 else f


@unittest.skipUnless(MASTER.exists(), '需要车队总表文件才能做这项核对')
class TestPasteIntoMaster(LiveServerCase):

    def test_block_lands_in_the_right_columns(self):
        # 1) 真实服务上走一遍
        mid = self.new_member('张三', '13800000000', is_owner=True)
        rid = self.new_round('十月采购')
        self.add_item(rid, mid, name='充电机', spec='3.3kwh', unit='个', price=6200,
                      purpose='高压电池充电', url='浩泰科技', paid=6200)
        self.add_item(rid, mid, name='绝缘手套', spec='/', unit='个', qty=2, price=250,
                      purpose='箱检', url='速迫')            # 没填实付
        items = self.detail(rid)['items']
        self.set_status(items[0]['id'], order_status='已到货')

        # 2) 取行块
        r = self.client.get('/api/rounds/%d/block' % rid)
        self.assertEqual(r.status_code, 200, r.text)
        payload = r.json()
        rows = [line.split('\t') for line in payload['tsv'].split('\n')]
        self.assertEqual(len(rows), 2, '两条条目应当就是两行')
        self.assertTrue(all(len(row) == 20 for row in rows), '每行必须是 20 列')

        # 3) 贴进总表副本（不碰原文件）
        tmp = Path(self._tmp.name) / 'master_copy.xlsx'
        patched_master_copy(tmp)
        wb = load_workbook(str(tmp))
        self.assertIn(TARGET_SHEET, wb.sheetnames)
        ws = wb[TARGET_SHEET]
        self.assertIn('组别', header_map(ws), '目标月份应当是含「组别」列的新表头')

        start = ws.max_row + 2          # 假装贴在「电气组」那一块的末尾
        for i, row in enumerate(rows):
            for j, val in enumerate(row, start=1):
                if val != '':
                    ws.cell(row=start + i, column=j, value=coerce(val))
        out = Path(self._tmp.name) / 'master_after.xlsx'
        wb.save(str(out))

        # 4) 读回来，按总表自己的表头逐列核对
        ws2 = load_workbook(str(out))[TARGET_SHEET]
        cols = header_map(ws2)

        def read(r_, name):
            return ws2.cell(row=r_, column=cols[name]).value

        r1, r2 = start, start + 1
        # 第一条：全部字段都填了
        self.assertEqual(read(r1, '物料名称/品牌'), '充电机')
        self.assertEqual(read(r1, '规格型号'), '3.3kwh')
        self.assertEqual(read(r1, '申请数量'), 1)
        self.assertEqual(read(r1, '单位'), '个')
        self.assertEqual(read(r1, '预计单价'), 6200)
        self.assertEqual(read(r1, '小计金额'), 6200)
        self.assertEqual(read(r1, '运费'), 0)
        self.assertEqual(read(r1, '合计金额'), 6200)
        self.assertEqual(read(r1, '用途'), '高压电池充电')
        self.assertEqual(read(r1, '申购人'), '张三')
        self.assertEqual(read(r1, '网址链接'), '浩泰科技')
        self.assertEqual(read(r1, '组别'), '电气组', '组别必须自动落在「组别」列')
        self.assertEqual(read(r1, '实付金额'), 6200)
        self.assertEqual(read(r1, '是否到货'), '是', '到货状态应落在「是否到货」列')

        # 第二条：没填实付的应当留空，而不是串到别的列
        self.assertEqual(read(r2, '物料名称/品牌'), '绝缘手套')
        self.assertEqual(read(r2, '申请数量'), 2)
        self.assertEqual(read(r2, '合计金额'), 500)
        self.assertEqual(read(r2, '组别'), '电气组')
        self.assertIsNone(read(r2, '实付金额'))
        self.assertIn(read(r2, '网址链接'), ('速迫',))

    def test_two_consecutive_pastes_do_not_shift(self):
        """连贴两次不该错位：第二块永远从紧凑的下一行开始。"""
        mid = self.new_member('李四', '15100000000')
        rid = self.new_round('九月采购')
        self.add_item(rid, mid, name='第一件', qty=1, unit='个', price=10, purpose='测试')
        first = self.client.get('/api/rounds/%d/block' % rid).json()['tsv']
        self.add_item(rid, mid, name='第二件', qty=1, unit='个', price=20, purpose='测试')
        second = self.client.get('/api/rounds/%d/block' % rid).json()['tsv']

        first_rows = first.split('\n')
        second_rows = second.split('\n')
        self.assertEqual(len(first_rows), 1)
        self.assertEqual(len(second_rows), 2)
        self.assertEqual(second_rows[0], first_rows[0], '前一行的内容不该变')
        self.assertTrue(second_rows[1].startswith('第二件\t'))


if __name__ == '__main__':
    unittest.main(verbosity=2)
