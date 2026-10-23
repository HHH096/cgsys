# -*- coding: utf-8 -*-
"""P1 导出：按在线表格模板生成 `总表` / `低于一千` / `高于一千`。"""
import unittest

from app.core import exporter, repo
from tests.support import TempDataTest

MAIN_HEADERS = ['物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
                '合计金额', '库存', '用途', '申购人', '领用时间', '申购人电话', '样式/图片',
                '网址链接', '组别', '备注', '实付金额', '发票', '是否到货']

SPLIT_HEADERS = ['物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
                 '合计金额', '库存', '用途', '申购人', '领用时间', '实际金额', '申购人电话',
                 '样式/图片', '网址链接', '组别', '备注', '实际金额', '发票']

HIGH_SUPPLIER = '浩泰科技'      # 6200 -> 高于一千
LOW_SUPPLIER = '速迫'           # 2 × 250 = 500 -> 低于一千


def headers(ws, row=2):
    return [str(v).strip() for v in next(ws.iter_rows(min_row=row, max_row=row, values_only=True)) if v is not None]


class TestExport(TempDataTest):

    def setUp(self):
        super().setUp()
        self.mid = repo.add_member('张三', '13800000000')
        self.rid = repo.create_round('八月采购')
        self.high_id = repo.add_item(self.rid, self.mid, {
            'name': '充电机', 'spec': '3.3kwh', 'qty': 1, 'unit': '个', 'price': 6200,
            'freight': 0, 'purpose': '高压电池充电', 'applicant': '张三',
            'phone': '13800000000', 'url': HIGH_SUPPLIER, 'group_name': '电气组',
            'paid': 6200, 'invoice': '有', 'arrived': '是', 'has_contract': '有',
        })
        self.low_id = repo.add_item(self.rid, self.mid, {
            'name': '绝缘手套', 'spec': '/', 'qty': 2, 'unit': '个', 'price': 250,
            'freight': 0, 'purpose': '箱检', 'applicant': '张三',
            'phone': '13800000000', 'url': LOW_SUPPLIER, 'group_name': '电气组',
            'paid': 500, 'invoice': '', 'arrived': '是',
        })
        self.round = repo.get_round(self.rid)
        self.items = repo.list_items(self.rid)
        self.wb, self.detail = exporter.build_workbook(self.round, self.items)

    # ---------------------------------------------------------- 结构

    def test_sheet_names_and_order(self):
        self.assertEqual(self.wb.sheetnames, ['总表', '低于一千', '高于一千'])

    def test_main_sheet_has_20_columns(self):
        ws = self.wb['总表']
        self.assertEqual(headers(ws), MAIN_HEADERS)
        self.assertIsNone(ws.cell(row=2, column=21).value)

    def test_split_sheets_have_21_columns(self):
        low = headers(self.wb['低于一千'])
        high = headers(self.wb['高于一千'])
        self.assertEqual(low, SPLIT_HEADERS + ['是否到货'])
        self.assertEqual(high, SPLIT_HEADERS + ['是否有合同'])

    def test_titles(self):
        for sheet, span in (('总表', 'A1:O1'), ('低于一千', 'A1:P1'), ('高于一千', 'A1:P1')):
            ws = self.wb[sheet]
            self.assertEqual(ws['A1'].value, self.round['title'])
            self.assertIn(span, [str(r) for r in ws.merged_cells.ranges])

    # ---------------------------------------------------------- 两套列映射

    def test_supplier_goes_to_the_right_column_in_each_layout(self):
        """总表「网址链接」在第 15 列，分表在第 16 列（分表多插了一列实际金额）。"""
        self.assertEqual(self.wb['总表'].cell(row=3, column=15).value, HIGH_SUPPLIER)
        self.assertEqual(self.wb['低于一千'].cell(row=3, column=16).value, LOW_SUPPLIER)
        self.assertEqual(self.wb['高于一千'].cell(row=3, column=16).value, HIGH_SUPPLIER)

    def test_phone_column_shifts_by_one_in_split_sheets(self):
        self.assertEqual(self.wb['总表'].cell(row=2, column=13).value, '申购人电话')
        self.assertEqual(self.wb['总表'].cell(row=3, column=13).value, '13800000000')
        for sheet in ('低于一千', '高于一千'):
            ws = self.wb[sheet]
            self.assertEqual(ws.cell(row=2, column=13).value, '实际金额')
            self.assertEqual(ws.cell(row=2, column=14).value, '申购人电话')
            self.assertIsNone(ws.cell(row=3, column=13).value, '%s 第 13 列应当留空' % sheet)
            self.assertEqual(ws.cell(row=3, column=14).value, '13800000000')

    def test_paid_goes_to_column_18_in_main_and_19_in_split(self):
        self.assertEqual(self.wb['总表'].cell(row=2, column=18).value, '实付金额')
        self.assertEqual(self.wb['总表'].cell(row=3, column=18).value, 6200)
        for sheet in ('低于一千', '高于一千'):
            ws = self.wb[sheet]
            self.assertEqual(ws.cell(row=2, column=19).value, '实际金额')

    def test_high_sheet_last_column_is_has_contract(self):
        self.assertEqual(self.wb['高于一千'].cell(row=3, column=21).value, '有')
        self.assertEqual(self.wb['低于一千'].cell(row=3, column=21).value, '是')

    # ---------------------------------------------------------- 金额

    def test_amounts_are_written(self):
        ws = self.wb['总表']
        self.assertEqual(ws.cell(row=3, column=6).value, 6200.0)   # 小计金额
        self.assertEqual(ws.cell(row=3, column=8).value, 6200.0)   # 合计金额
        self.assertEqual(ws.cell(row=4, column=6).value, 500.0)
        self.assertEqual(ws.cell(row=4, column=8).value, 500.0)

    # ---------------------------------------------------------- 合计行

    def test_total_row_on_main_sheet(self):
        ws = self.wb['总表']
        r = 3 + len(self.items)
        self.assertEqual(ws.cell(row=r, column=1).value, '合计')
        self.assertEqual(ws.cell(row=r, column=8).value, '=SUM(H3:H4)')
        self.assertIn('A%d:G%d' % (r, r), [str(x) for x in ws.merged_cells.ranges])
        self.assertEqual(ws.cell(row=r, column=8).fill.fgColor.rgb, 'FFFFFF00')

    def test_total_row_label_in_column_g_on_low_sheet(self):
        ws = self.wb['低于一千']
        r = 4
        self.assertIsNone(ws.cell(row=r, column=1).value)
        self.assertEqual(ws.cell(row=r, column=7).value, '合计')
        self.assertEqual(ws.cell(row=r, column=8).value, '=SUM(H3:H3)')
        self.assertNotIn('A%d:G%d' % (r, r), [str(x) for x in ws.merged_cells.ranges])

    def test_total_row_merged_on_high_sheet(self):
        ws = self.wb['高于一千']
        r = 4
        self.assertEqual(ws.cell(row=r, column=1).value, '合计')
        self.assertEqual(ws.cell(row=r, column=8).value, '=SUM(H3:H3)')
        self.assertIn('A%d:G%d' % (r, r), [str(x) for x in ws.merged_cells.ranges])

    # ---------------------------------------------------------- 拆分

    def test_split_follows_the_supplier_rule(self):
        main = [ws.cell(row=3, column=1).value for ws in [self.wb['总表']]]
        self.assertEqual(main, ['充电机'])
        self.assertEqual(self.wb['总表'].cell(row=4, column=1).value, '绝缘手套')
        self.assertEqual(self.wb['高于一千'].cell(row=3, column=1).value, '充电机')
        self.assertEqual(self.wb['低于一千'].cell(row=3, column=1).value, '绝缘手套')

    def test_detail_counts_and_sums(self):
        self.assertEqual(self.detail['high'], 1)
        self.assertEqual(self.detail['low'], 1)
        self.assertEqual(self.detail['sum_high'], 6200.0)
        self.assertEqual(self.detail['sum_low'], 500.0)
        self.assertEqual(self.detail['sum_all'], 6700.0)

    def test_empty_round_does_not_crash(self):
        rid = repo.create_round('空轮次')
        wb, detail = exporter.build_workbook(repo.get_round(rid), [])
        self.assertEqual(wb.sheetnames, ['总表', '低于一千', '高于一千'])
        self.assertEqual(detail['sum_all'], 0.0)


class TestExportToFile(TempDataTest):

    def test_save_and_reload(self):
        mid = repo.add_member('张三')
        rid = repo.create_round('九月采购')
        repo.add_item(rid, mid, {'name': '强芯毡', 'qty': 4, 'unit': '张', 'price': 86.6,
                                 'freight': 10, 'purpose': '制作空套', 'applicant': '张三',
                                 'url': 'https://e.tb.cn/h.8RE', 'group_name': '电气组'})
        from app import config
        path = config.EXPORT_DIR / 'test.xlsx'
        config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        out, detail = exporter.save_round_xlsx(repo.get_round(rid), repo.list_items(rid), path)
        self.assertTrue(out.exists())

        from openpyxl import load_workbook
        wb = load_workbook(str(out))
        ws = wb['总表']
        self.assertEqual(ws.cell(row=3, column=1).value, '强芯毡')
        self.assertEqual(ws.cell(row=3, column=6).value, 346.4)     # 4 × 86.6
        self.assertEqual(ws.cell(row=3, column=8).value, 356.4)     # + 运费 10
        self.assertEqual(detail['sum_all'], 356.4)


class TestExportReflectsStatus(TempDataTest):
    """导出的「发票」「是否到货」两列跟着下单/发票状态走。"""

    def setUp(self):
        super().setUp()
        self.mid = repo.add_member('张三', '13800000000')
        self.rid = repo.create_round('八月采购')
        self.high = repo.add_item(self.rid, self.mid, {
            'name': '充电机', 'qty': 1, 'unit': '个', 'price': 6200,
            'purpose': '高压电池充电', 'applicant': '张三',
            'url': '浩泰科技', 'paid': 6200})
        self.low = repo.add_item(self.rid, self.mid, {
            'name': '绝缘手套', 'qty': 2, 'unit': '个', 'price': 250,
            'purpose': '箱检', 'applicant': '张三',
            'url': '速迫', 'paid': 500})

    def export(self):
        return exporter.build_workbook(repo.get_round(self.rid), repo.list_items(self.rid))

    def test_columns_are_empty_before_any_progress(self):
        wb, _ = self.export()
        main = wb['总表']
        self.assertIsNone(main.cell(row=3, column=19).value)     # 发票
        self.assertIsNone(main.cell(row=3, column=20).value)     # 是否到货

    def test_columns_follow_status(self):
        repo.set_item_status(self.high, order_status='已到货', invoice_status='已收到')
        repo.set_item_status(self.low, invoice_status='已要')     # 只是要了，还没拿到
        wb, _ = self.export()
        main = wb['总表']
        self.assertEqual(main.cell(row=3, column=19).value, '有')
        self.assertEqual(main.cell(row=3, column=20).value, '是')
        self.assertIsNone(main.cell(row=4, column=19).value)
        self.assertIsNone(main.cell(row=4, column=20).value)

    def test_split_sheets_carry_the_status_too(self):
        repo.set_item_status(self.high, order_status='已到货', invoice_status='已收到')
        wb, _ = self.export()
        high = wb['高于一千']
        self.assertEqual(high.cell(row=2, column=20).value, '发票')
        self.assertEqual(high.cell(row=3, column=20).value, '有')

        low = wb['低于一千']
        self.assertEqual(low.cell(row=3, column=20).value, None)
        self.assertEqual(low.cell(row=2, column=21).value, '是否到货')

    def test_has_contract_lands_in_the_high_sheet_last_column(self):
        repo.update_item(self.high, {'has_contract': '有'})
        wb, _ = self.export()
        self.assertEqual(wb['高于一千'].cell(row=3, column=21).value, '有')

    def test_status_change_survives_re_export(self):
        """改完状态再导一次，结果要跟上（不是只改内存里的那条记录）。"""
        repo.set_item_status(self.high, order_status='已到货')
        wb1, _ = self.export()
        self.assertEqual(wb1['总表'].cell(row=3, column=20).value, '是')

        repo.set_item_status(self.high, order_status='已下单')
        wb2, _ = self.export()
        self.assertIsNone(wb2['总表'].cell(row=3, column=20).value)


if __name__ == '__main__':
    unittest.main(verbosity=2)


class TestFormulaInjection(unittest.TestCase):
    """导出时**必须防公式注入**。

    openpyxl 会把 `=` 开头的字符串当公式写进单元格。组员在物料名称或备注里
    填 `=HYPERLINK("http://坏站","点我")`，采购一打开导出的表格就可能被带到外链。
    前面加单引号，Excel 按普通文本处理。
    """

    def test_an_equals_prefix_is_escaped(self):
        got = exporter._blank('=HYPERLINK("http://x","y")')
        self.assertEqual(got, '\'=HYPERLINK("http://x","y")')

    def test_plus_minus_at_are_escaped_too(self):
        for s in ('+1+1', '-1+1', '@SUM(A1)'):
            self.assertTrue(exporter._blank(s).startswith("'"), s)

    def test_normal_text_is_untouched(self):
        self.assertEqual(exporter._blank('电容 22uF'), '电容 22uF')

    def test_blank_still_becomes_none(self):
        self.assertIsNone(exporter._blank('   '))
        self.assertIsNone(exporter._blank(None))

    def test_the_block_export_escapes_as_well(self):
        from app.core import blockexport
        self.assertEqual(blockexport._safe('=1+1'), "'=1+1")
        self.assertEqual(blockexport._safe('排线'), '排线')
        self.assertIsNone(blockexport._safe(''))
