# -*- coding: utf-8 -*-
"""P4 行块接口：/api/rounds/{id}/block 与 /block.xlsx。"""
import io
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.server.api import app
from tests.support import TempDataTest, purchase_client

HEADERS = [
    '物料名称/品牌', '规格型号', '申请数量', '单位', '预计单价', '小计金额', '运费',
    '合计金额', '库存', '用途', '申购人', '领用时间', '申购人电话', '样式/图片',
    '网址链接', '组别', '备注', '实付金额', '发票', '是否到货',
]


class BlockCase(TempDataTest):

    def setUp(self):
        super().setUp()
        self.client = purchase_client(self)
        self.rid = self.client.post('/api/rounds', json={'name': '八月采购'}).json()['id']
        self.mid = self.client.post('/api/members', json={'name': '张三'}).json()['id']

    def add(self, name, **over):
        body = {'member_id': self.mid, 'name': name, 'unit': '个', 'qty': 1,
                'purpose': '用', 'price': 100, 'freight': 0,
                'applicant': '张三', 'group_name': '电气组'}
        body.update(over)
        r = self.client.post('/api/rounds/%d/items' % self.rid, json=body)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['item']

    def block(self, rid=None):
        r = self.client.get('/api/rounds/%d/block' % (rid or self.rid))
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()


class TestBlockJson(BlockCase):

    def test_headers_and_group(self):
        j = self.block()
        self.assertEqual(j['headers'], HEADERS)
        self.assertEqual(j['group'], '电气组')

    def test_empty_round_is_ok_not_an_error(self):
        j = self.block()
        self.assertEqual(j['rows'], [])
        self.assertEqual(j['tsv'], '')
        self.assertEqual(j['reconcile']['count'], 0)

    def test_missing_round_is_404(self):
        self.assertEqual(self.client.get('/api/rounds/9999/block').status_code, 404)

    def test_rows_are_positioned_like_the_master_table(self):
        self.add('充电机', spec='3.3kwh', price=6200, paid=6200, url='浩泰科技')
        self.add('绝缘手套', spec='/', qty=2, price=250, url='速迫')
        j = self.block()
        self.assertEqual(len(j['rows']), 2)

        row = j['rows'][0]
        self.assertEqual(row[0], '充电机')
        self.assertEqual(row[1], '3.3kwh')
        self.assertEqual(row[7], 6200)          # 合计金额
        self.assertEqual(row[9], '用')           # 用途
        self.assertEqual(row[14], '浩泰科技')     # 网址链接
        self.assertEqual(row[15], '电气组')       # 组别（自动填）
        self.assertEqual(row[17], 6200)         # 实付金额

    def test_tsv_has_no_header_row(self):
        self.add('充电机', price=6200)
        self.add('绝缘手套', qty=2, price=250)
        tsv = self.block()['tsv']
        lines = tsv.split('\n')
        self.assertEqual(len(lines), 2, 'TSV 不得含表头行')
        self.assertTrue(lines[0].startswith('充电机\t'))
        self.assertNotIn('物料名称/品牌', tsv)
        self.assertEqual(len(lines[0].split('\t')), 20)

    def test_reconcile_numbers(self):
        self.add('充电机', price=6200, paid=6200, url='浩泰科技')
        self.add('绝缘手套', qty=2, price=250, url='速迫')      # 没填实付
        rec = self.block()['reconcile']
        self.assertEqual(rec['count'], 2)
        self.assertEqual(rec['estimate'], 6700.0)
        self.assertEqual(rec['paid'], 6200.0)
        self.assertEqual(rec['diff'], -500.0)
        self.assertEqual(rec['paid_count'], 1)
        self.assertEqual(rec['unpriced'], 1)
        self.assertEqual([s['supplier'] for s in rec['suppliers']], ['浩泰科技', '速迫'])

    def test_note_warns_that_old_months_differ(self):
        self.assertIn('7月下旬起', self.block()['note'])


class TestBlockXlsx(BlockCase):

    def test_downloads_a_workbook(self):
        self.add('充电机', price=6200, paid=6200)
        self.add('绝缘手套', qty=2, price=250)
        r = self.client.get('/api/rounds/%d/block.xlsx' % self.rid)
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.headers['content-type'])
        self.assertEqual(r.content[:2], b'PK')

        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['行块', '对账'])
        ws = wb['行块']
        self.assertEqual(ws.cell(row=1, column=1).value, '充电机',
                         '第一行就该是数据，不能是表头')
        self.assertEqual(ws.cell(row=2, column=1).value, '绝缘手套')
        self.assertEqual(ws.cell(row=1, column=16).value, '电气组')

    def test_reconcile_sheet(self):
        self.add('充电机', price=6200, paid=6200)
        r = self.client.get('/api/rounds/%d/block.xlsx' % self.rid)
        ws = load_workbook(io.BytesIO(r.content))['对账']
        rows = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value
                for i in range(1, ws.max_row + 1)}
        self.assertEqual(rows['条目数'], 1)
        self.assertEqual(rows['预计合计（合计金额之和）'], 6200)

    def test_empty_round_is_400(self):
        self.assertEqual(self.client.get('/api/rounds/%d/block.xlsx' % self.rid).status_code, 400)

    def test_missing_round_is_404(self):
        self.assertEqual(self.client.get('/api/rounds/9999/block.xlsx').status_code, 404)

    def test_saved_under_export_dir(self):
        from app import config
        self.add('充电机', price=6200)
        self.client.get('/api/rounds/%d/block.xlsx' % self.rid)
        self.assertEqual(len(list(config.EXPORT_DIR.glob('行块_*.xlsx'))), 1)


class TestFrontendWiring(unittest.TestCase):
    """回填卡片确实接上了（防止以后误删）。"""

    @classmethod
    def setUpClass(cls):
        root = Path(__file__).resolve().parent.parent
        cls.html = (root / 'app/web/index.html').read_text(encoding='utf-8')
        cls.js = (root / 'app/web/app.js').read_text(encoding='utf-8')

    def test_controls_exist(self):
        for needle in ('id="blockText"', 'id="blockCopy"', 'id="blockXlsx"',
                       'id="blockRecon"', 'id="blockHint"', 'id="blockCols"'):
            self.assertIn(needle, self.html, 'index.html 缺 %s' % needle)

    def test_textarea_is_readonly(self):
        idx = self.html.index('id="blockText"')
        self.assertIn('readonly', self.html[idx:idx + 120],
                      '行块框应当是只读的，避免误改后再复制')

    def test_buttons_are_bound(self):
        self.assertIn("$('blockCopy').addEventListener('click', blockCopy)", self.js)
        self.assertIn("$('blockXlsx').addEventListener('click', blockXlsx)", self.js)

    def test_points_at_both_endpoints(self):
        self.assertIn("'/block'", self.js)
        self.assertIn("'/block.xlsx'", self.js)

    def test_rendered_on_load_and_refresh(self):
        self.assertIn('renderBlock();', self.js, 'renderAll 里要刷新行块')
        # 换轮次时要重新取行块；三个采购请求现在并行发，所以断言的是并行写法
        self.assertIn('Promise.all([loadBlock(),', self.js, '换轮次时要重新取行块')

    def test_copy_falls_back_when_clipboard_api_is_unavailable(self):
        """局域网走 http，navigator.clipboard 会被浏览器拒绝，必须有降级。"""
        self.assertIn("execCommand('copy')", self.js)
        self.assertIn('Ctrl+C', self.js)


if __name__ == '__main__':
    unittest.main(verbosity=2)
